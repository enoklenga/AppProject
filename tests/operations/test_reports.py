from datetime import date
from decimal import Decimal
from io import BytesIO

from django.contrib.auth import get_user_model
from django.test import TestCase
from django.urls import reverse
from openpyxl import load_workbook

from apps.operations.report_services import FiltriReport, crea_report_xlsx
from apps.projects.models import (
    Assegnazione,
    Cliente,
    Commessa,
    TariffaAssegnazione,
)
from apps.timesheets.models import RigaOre, SpesaTrasferta

User = get_user_model()


class ReportTests(TestCase):
    def setUp(self):
        self.admin = User.objects.create_user(
            email="admin-test@example.com",
            password="Password-test-123",
            ruolo=User.Ruolo.ADMIN,
            is_staff=True,
        )
        self.consulente = User.objects.create_user(
            email="consulente-test@example.com",
            password="Password-test-123",
        )
        self.cliente = Cliente.objects.create(
            ragione_sociale="Cliente Report",
            partita_iva="IT00000000081",
        )
        self.commessa = Commessa.objects.create(
            cliente=self.cliente,
            codice="REP-001",
            descrizione="Report mensile",
            data_inizio=date(2026, 1, 1),
        )
        self.assegnazione = Assegnazione.objects.create(
            consulente=self.consulente,
            commessa=self.commessa,
            ore_previste=80,
            data_inizio=date(2026, 1, 1),
        )
        TariffaAssegnazione.objects.create(
            assegnazione=self.assegnazione,
            tipo_attivita=TariffaAssegnazione.TipoAttivita.CONSULENZA,
            tariffa_oraria=Decimal("100.00"),
            valida_dal=date(2026, 1, 1),
            creata_da=self.admin,
        )
        RigaOre.objects.create(
            assegnazione=self.assegnazione,
            data=date(2026, 7, 10),
            tipo_attivita=TariffaAssegnazione.TipoAttivita.CONSULENZA,
            ore=4,
            nota="Attività report",
            inserita_da=self.consulente,
            ultima_modifica_da=self.consulente,
        )
        SpesaTrasferta.objects.create(
            assegnazione=self.assegnazione,
            data=date(2026, 7, 10),
            categoria=SpesaTrasferta.Categoria.VIAGGIO,
            importo=Decimal("25.00"),
            nota="Pedaggio",
            inserita_da=self.consulente,
            ultima_modifica_da=self.consulente,
        )

    def test_excel_contiene_fogli_e_totali(self):
        contenuto = crea_report_xlsx(FiltriReport(anno=2026, mese=7))
        workbook = load_workbook(BytesIO(contenuto), data_only=True)

        self.assertEqual(
            workbook.sheetnames,
            [
                "Riepilogo",
                "Ore",
                "Spese",
                "Per cliente",
                "Per commessa",
                "Per fase",
                "Per consulente",
            ],
        )
        self.assertEqual(workbook["Riepilogo"]["B3"].value, 4)
        self.assertEqual(workbook["Riepilogo"]["B6"].value, 425)
        self.assertEqual(workbook["Ore"].max_row, 2)
        self.assertEqual(workbook["Spese"].max_row, 2)


    def test_excel_non_interpreta_note_utente_come_formule(self):
        riga = RigaOre.objects.get()
        riga.nota = "=HYPERLINK(\"https://example.invalid\",\"click\")"
        riga.save(update_fields=["nota"])

        contenuto = crea_report_xlsx(FiltriReport(anno=2026, mese=7))
        workbook = load_workbook(BytesIO(contenuto), data_only=False)
        cella = workbook["Ore"]["J2"]

        self.assertTrue(str(cella.value).startswith("'="))
        self.assertEqual(cella.data_type, "s")

    def test_admin_puo_scaricare_excel(self):
        self.client.force_login(self.admin)
        response = self.client.get(
            reverse("operations:report-mensile-xlsx"),
            {"mese": "2026-07"},
        )
        self.assertEqual(response.status_code, 200)
        self.assertIn(
            "spreadsheetml",
            response["Content-Type"],
        )

    def test_consulente_non_puo_accedere(self):
        self.client.force_login(self.consulente)
        response = self.client.get(
            reverse("operations:report-mensile"),
            {"mese": "2026-07"},
        )
        self.assertEqual(response.status_code, 403)
